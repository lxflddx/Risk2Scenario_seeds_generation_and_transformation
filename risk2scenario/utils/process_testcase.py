import os
import json
from openpyxl import Workbook
"""
将提取的逻辑场景转化成表格，方便后续仿真
"""
# ================= 配置 =================
# True：数字文件名自动转换为 TC_01 格式
# False：保留原文件名
USE_TC_FORMAT = True

# 输入文件夹路径
folder_path = r"D:\pythonProject\LLM_Scenario\data\my_convert_method_results\SR"

# 输出 Excel 文件
excel_path = r"D:\pythonProject\LLM_Scenario\data\my_convert_method_results\SR.xlsx"

# ========================================

test_list = []

for filename in os.listdir(folder_path):
    if not filename.endswith(".json"):
        continue

    file_path = os.path.join(folder_path, filename)
    raw_id = os.path.splitext(filename)[0]

    # 处理 testcase_id
    if USE_TC_FORMAT:
        if not raw_id.isdigit():
            print(f"Skip invalid numeric filename: {filename}")
            continue
        testcase_id = f"TC_{int(raw_id):02d}"
    else:
        testcase_id = raw_id

    # 读取 JSON 文件
    with open(file_path, "r", encoding="utf-8") as file:
        try:
            data = json.load(file)
            logical_scenario = data.get("logical_scenario")

            if not logical_scenario:
                continue

            # ========= 字符串型 =========
            if isinstance(logical_scenario, str):
                # 去掉多余空行
                lines = [line.rstrip() for line in logical_scenario.splitlines() if line.strip()]
                # 第一行保持原样，后续行缩进 4 个空格
                if lines:
                    formatted = [lines[0]] + ["    " + line for line in lines[1:]]
                    scenario_content = "\n".join(formatted)
                else:
                    scenario_content = ""

            # ========= 字典型 =========
            elif isinstance(logical_scenario, dict):
                scenario_content = json.dumps(logical_scenario, ensure_ascii=False, indent=2)

            else:
                print(f"Unsupported logical_scenario type in {filename}")
                continue

            test_list.append((testcase_id, scenario_content))

        except Exception as e:
            print(f"Error processing {filename}: {e}")

# 排序
if USE_TC_FORMAT:
    test_list.sort(key=lambda x: int(x[0].split("_")[1]))
else:
    test_list.sort(key=lambda x: x[0].lower())

# 保存到 Excel
wb = Workbook()
ws = wb.active
ws.title = "Test Cases"

ws.cell(row=1, column=1, value="Testcase ID")
ws.cell(row=1, column=2, value="Logical Scenario")

for i, (testcase_id, scenario) in enumerate(test_list, start=2):
    ws.cell(row=i, column=1, value=testcase_id)
    ws.cell(row=i, column=2, value=scenario)

wb.save(excel_path)

print(f"Saved to {excel_path}, total {len(test_list)} testcases")
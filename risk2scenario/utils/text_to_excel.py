import os
import openpyxl
from openpyxl import Workbook

"""
读取一个文件中的txt文件，然后将其中的内容写入到一个excel文件中。用于处理事故报告或者事故描述，把它们存储到Excel文件中，用于后续转换成测试用例。
"""

def txt_to_excel(input_folder, output_file):
    # 创建一个新的Excel工作簿
    wb = Workbook()
    ws = wb.active

    # 添加表头
    ws['A1'] = 'Case ID'
    ws['B1'] = 'Content'

    # 遍历文件夹中的所有txt文件
    row = 2  # 从第二行开始写入数据
    for filename in os.listdir(input_folder):
        if filename.endswith('.txt'):
            filepath = os.path.join(input_folder, filename)

            # 读取txt文件内容
            with open(filepath, 'r', encoding='utf-8') as file:
                content = file.read()

            case_id = os.path.splitext(filename)[0]
            ws.cell(row=row, column=1, value=case_id)  # Case ID列写入不带后缀的文件名
            ws.cell(row=row, column=2, value=content)  # Content列写入文件内容
            row += 1

    # 保存Excel文件
    wb.save(output_file)
    print(f"数据已成功写入 {output_file}")


# 使用示例
input_folder = 'D:\\00LLMScenario\\dataset\\highway_gemma'  # 存放txt文件的文件夹路径
output_file = '../data/my_final_straight_new.xlsx'  # 输出的Excel文件名
txt_to_excel(input_folder, output_file)
